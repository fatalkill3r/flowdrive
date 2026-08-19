import csv, io, json, subprocess
from .metadata import IMAGES, AUDIO, VIDEOS
from .storage import open_file
TEXT={"txt","log","csv","json","xml","md","sh","py"}; SPREADSHEETS={"xls","xlsx"}; WORD={"doc","docx"}; MAX=1024*1024; MAX_DOCUMENT=50*1024*1024; MAX_ROWS=200; MAX_COLS=50; MAX_SHEETS=10
class PreviewError(Exception):pass
def preview_kind(item):
    ext=item.extension.lower()
    if ext in IMAGES:return "image"
    if ext=="pdf":return "pdf"
    if ext in AUDIO:return "audio"
    if ext in VIDEOS:return "video"
    if ext in SPREADSHEETS:return "spreadsheet"
    if ext in WORD:return "document"
    if ext in TEXT:return ext if ext in ("csv","json") else "text"
    return "unsupported"
def text_content(item):
    with open_file(item) as stream: raw=stream.read(MAX)
    text=raw.decode("utf-8",errors="replace")
    if item.extension=="json":
        try:return json.dumps(json.loads(text),indent=2)
        except ValueError:return text
    return text
def csv_rows(item):
    return [row[:MAX_COLS] for row in csv.reader(io.StringIO(text_content(item)))][:MAX_ROWS]
def spreadsheet_sheets(item):
    if item.size>MAX_DOCUMENT:raise PreviewError("This spreadsheet is too large to preview. Download it to open locally.")
    try:
        with open_file(item) as stream:
            if item.extension.lower()=="xlsx":
                from openpyxl import load_workbook
                book=load_workbook(stream,read_only=True,data_only=True)
                sheets=[]
                try:
                    for sheet in book.worksheets[:MAX_SHEETS]:
                        rows=[["" if value is None else value for value in row[:MAX_COLS]] for row in sheet.iter_rows(max_row=MAX_ROWS,values_only=True)]
                        sheets.append({"name":sheet.title,"rows":rows})
                finally:book.close()
                return sheets
            import xlrd
            book=xlrd.open_workbook(file_contents=stream.read())
            return [{"name":sheet.name,"rows":[sheet.row_values(index,0,min(sheet.ncols,MAX_COLS)) for index in range(min(sheet.nrows,MAX_ROWS))]} for sheet in book.sheets()[:MAX_SHEETS]]
    except PreviewError:raise
    except Exception as exc:raise PreviewError("This spreadsheet could not be read. It may be damaged or password-protected.") from exc
def document_content(item):
    if item.size>MAX_DOCUMENT:raise PreviewError("This document is too large to preview. Download it to open locally.")
    try:
        with open_file(item) as stream:
            if item.extension.lower()=="docx":
                from docx import Document
                document=Document(stream); parts=[paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
                return "\n".join(parts)[:MAX]
            result=subprocess.run(["antiword",stream.name],capture_output=True,timeout=15,check=False)
            if result.returncode:return_code_error(result.stderr)
            return result.stdout[:MAX].decode("utf-8",errors="replace")
    except PreviewError:raise
    except Exception as exc:raise PreviewError("This Word document could not be read. It may be damaged or password-protected.") from exc
def return_code_error(message):
    raise PreviewError((message or b"This Word document could not be read.").decode("utf-8",errors="replace")[:240])
